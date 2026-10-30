import sqlite3
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from araco import backup, config, exporter, importer, ocr
from araco.context import NotAllowed
from araco.services import customers, tickets


def test_import_xlsx_with_arabic_headers_merges_sites(admin, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["الاسم", "الشركة", "الهاتف", "المنطقة", "الورشة", "ملاحظات"])
    ws.append(["جورج حداد", "حداد للمقاولات", 3123456, "زحلة", "بناية الكرك", ""])
    ws.append(["جورج حداد", "حداد للمقاولات", 3123456, "زحلة", "فيلا كسارة", ""])
    ws.append(["Rami Saad", "", "70 111 222", "Jbeil", "", "cash"])
    ws.append(["", "", "", "", "", "orphan note"])
    path = tmp_path / "old.xlsx"
    wb.save(path)

    header, rows = importer.read_table(path)
    mapping = importer.guess_mapping(header)
    assert mapping == {"name": 0, "company": 1, "phone": 2, "area": 3, "site": 4, "notes": 5}
    stats = importer.import_customers(admin, rows, mapping)
    assert (stats["created"], stats["sites"], len(stats["errors"])) == (2, 2, 1)
    george = customers.search(admin.db, "جورج")[0]
    assert george["phone"] == "3123456" and {s["name"] for s in customers.sites(admin.db, george["id"])} == \
        {"بناية الكرك", "فيلا كسارة"}
    # importing again does not duplicate
    again = importer.import_customers(admin, rows, mapping)
    assert again["created"] == 0 and again["skipped"] == 2


def test_import_csv_cp1256_and_role(admin, entry, tmp_path):
    path = tmp_path / "old.csv"
    path.write_bytes("Code;Customer Name;Phone\nK-7;سمير;01 222333\n".encode("cp1256"))
    header, rows = importer.read_table(path)
    mapping = importer.guess_mapping(header)
    assert mapping == {"code": 0, "name": 1, "phone": 2}
    with pytest.raises(NotAllowed):
        importer.import_customers(entry, rows, mapping)
    importer.import_customers(admin, rows, mapping)
    assert customers.find_by_code(admin.db, "K-7")["name"] == "سمير"


def test_export_xlsx_and_pdf(tmp_path):
    headers = ["رقم البون", "الزبون", "الكمية"]
    rows = [["1001", "شركة الأمل للمقاولات", 8.0], ["1002", "Beirut Builders", 12.5]]
    x = exporter.export_xlsx(tmp_path / "r.xlsx", "تقرير", headers, rows, totals=["", "المجموع", 20.5], rtl=True)
    ws = load_workbook(x).active
    assert ws.sheet_view.rightToLeft and ws["B5"].value == "شركة الأمل للمقاولات" and ws["C7"].value == 20.5
    p = exporter.export_pdf(tmp_path / "r.pdf", "تقرير", headers, rows, totals=["", "المجموع", 20.5], rtl=True)
    assert p.read_bytes()[:4] == b"%PDF"
    assert exporter.shape("abc") == "abc" and exporter.shape("سلام") != "سلام"


def test_daily_backup_keeps_30_and_restore(admin, tmp_path):
    start = date(2026, 1, 1)
    for i in range(35):
        backup.daily_backup(admin.db, today=start + timedelta(days=i))
    autos = sorted(p.name for p in config.backups_dir().glob("auto_*.db"))
    assert len(autos) == 30 and autos[0] == "auto_2026-01-06.db"
    assert backup.daily_backup(admin.db, today=start + timedelta(days=34)) is None  # already done today

    cid = customers.create(admin, {"name": "Before backup"})
    snap = backup.backup_now(admin.db)
    customers.create(admin, {"name": "After backup"})
    zipped = backup.full_zip(admin.db, tmp_path)
    assert backup.is_araco_db(snap) and not backup.is_araco_db(tmp_path / "r.none")
    safety = backup.restore(admin.db, snap)
    assert [c["name"] for c in customers.search(admin.db)] == ["Before backup"]
    assert customers.get(admin.db, cid)
    assert sqlite3.connect(safety).execute("SELECT COUNT(*) FROM customers").fetchone()[0] == 2
    backup.restore(admin.db, zipped)
    assert len(customers.search(admin.db)) == 2


def test_ocr_field_extraction():
    text = "ARACO READY MIX\nDelivery Ticket No: 458213\nDate: 14/10/2026  Tel 03 123456\nB300  Qty 9.5 m3"
    assert ocr.extract_fields(text) == {"ticket_no": "458213", "date": "2026-10-14", "quantity": 9.5}
    ar = "رقم البون: ٤٥٦٧٨\nالتاريخ ٢٠٢٦/١٠/٠٣\nالكمية ١٢ م3"
    assert ocr.extract_fields(ar) == {"ticket_no": "45678", "date": "2026-10-03", "quantity": 12.0}
    assert ocr.extract_fields("nothing useful here") == {}


@pytest.mark.skipif(not ocr.find_tesseract(), reason="tesseract not installed")
def test_ocr_reads_generated_ticket(tmp_path):
    from araco.seed import draw_ticket_image
    path = draw_ticket_image(tmp_path / "t.jpg", "772310", "2026-10-21", "08:15", "Client", "B350", 7.5, "M-14")
    fields = ocr.extract_fields(ocr.read_text(path, ocr.find_tesseract()))
    assert fields.get("ticket_no") == "772310"
    assert fields.get("date") == "2026-10-21"
    assert fields.get("quantity") == 7.5


def test_scan_pdf_attachment_renders(admin, tmp_path):
    from araco import media
    pdf = exporter.export_pdf(tmp_path / "scan.pdf", "Bon 55", ["a"], [["b"]])
    pages = media.load_pages(pdf)
    assert len(pages) == 1 and pages[0].width > 500
    cid = customers.create(admin, {"name": "X"})
    tid = tickets.create(admin, {"ticket_no": "55", "date": "2026-10-10", "customer_id": cid, "grade": "B250",
                                 "quantity": 3}, scan_files=[pdf])
    assert tickets.scans(admin.db, tid)[0]["file_path"].endswith(".pdf")


def test_seed_and_wipe(admin):
    from araco import seed
    from araco.services import reports
    stats = seed.seed(admin, n_tickets=60, with_images=3, today=date(2026, 10, 30))
    assert stats == {"customers": 20, "tickets": 60, "images": 3}
    rows, total = reports.summary(admin.db, "customer")
    assert total["count"] > 50 and len(rows) <= 20
    seed.wipe_business_data(admin)
    assert customers.search(admin.db) == [] and tickets.search(admin.db) == []
