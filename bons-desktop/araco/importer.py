"""Import historical customers and tickets from Excel (.xlsx) or CSV with a column mapping."""
import csv
import io
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from .context import ValidationError
from .db import digits, norm
from .services import audit, customers

TARGETS = ["code", "name", "company", "phone", "phone2", "address", "area", "payment_terms", "notes", "site"]

# Header words (normalised) that suggest each target field.
SYNONYMS = {
    "code": ["code", "customer code", "client code", "id", "no", "number", "كود", "الكود", "كود الزبون", "رمز", "رقم الزبون", "الرقم"],
    "name": ["name", "customer", "client", "customer name", "client name", "الاسم", "اسم", "الزبون", "اسم الزبون",
             "العميل", "اسم العميل"],
    "company": ["company", "firm", "contractor", "الشركه", "شركه", "المتعهد", "المؤسسه"],
    "phone": ["phone", "tel", "telephone", "mobile", "gsm", "phone 1", "الهاتف", "هاتف", "الخليوي", "رقم الهاتف",
              "جوال", "موبايل"],
    "phone2": ["phone 2", "phone2", "tel 2", "mobile 2", "other phone", "هاتف 2", "هاتف اخر"],
    "address": ["address", "location", "العنوان", "عنوان"],
    "area": ["area", "region", "city", "town", "المنطقه", "منطقه", "البلده", "المدينه"],
    "payment_terms": ["payment", "payment terms", "terms", "شروط الدفع", "الدفع", "طريقه الدفع"],
    "notes": ["notes", "note", "remarks", "comments", "ملاحظات", "ملاحظه"],
    "site": ["site", "sites", "project", "project site", "الورشه", "ورشه", "الموقع", "المشروع", "المشاريع"],
}


def read_table(path) -> tuple[list[str], list[list[str]]]:
    """Return (header, rows) as strings. The first row is the header."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1256", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    elif ext == ".xls":
        raise ValidationError(["err_xls_old"])
    else:
        raise ValidationError(["err_file_type"])
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return rows[0], rows[1:]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # phone numbers stored as numbers in Excel
    return str(v).strip()


def guess_mapping(header: list[str], synonyms: dict | None = None) -> dict:
    """{target: column index} from header names (Arabic or English)."""
    synonyms = {t: [norm(w) for w in words] for t, words in (synonyms or SYNONYMS).items()}
    mapping = {}
    for idx, h in enumerate(header):
        key = re.sub(r"[_\-.:#]+", " ", norm(h)).strip()
        for target, words in synonyms.items():
            if target in mapping:
                continue
            if key in words:
                mapping[target] = idx
                break
    # loose match on contained words for anything left
    for idx, h in enumerate(header):
        key = norm(h)
        if idx in mapping.values():
            continue
        for target, words in synonyms.items():
            if target not in mapping and any(len(w) > 3 and w in key for w in words):
                mapping[target] = idx
                break
    return mapping


def import_customers(ctx, rows: list[list[str]], mapping: dict, update_existing: bool = False) -> dict:
    """Create customers from mapped rows.

    Rows describing the same customer (same code, or same name + phone) are merged, so a sheet
    with one row per project site creates one customer with several sites.
    Existing customers (same code, or same name + phone) are skipped, or updated if asked.
    """
    ctx.require_admin()
    if "name" not in mapping:
        raise ValidationError(["err_map_name"])
    stats = {"created": 0, "updated": 0, "skipped": 0, "sites": 0, "errors": []}
    groups: dict[str, dict] = {}
    for line_no, row in enumerate(rows, start=2):
        rec = {t: (row[i] if i is not None and i < len(row) else "").strip() for t, i in mapping.items()}
        if not rec.get("name"):
            if any(rec.values()):
                stats["errors"].append((line_no, "err_name_required"))
            continue
        key = norm(rec.get("code")) or f"{norm(rec['name'])}|{digits(rec.get('phone'))}"
        g = groups.setdefault(key, {"data": {}, "sites": []})
        for t, v in rec.items():
            if t == "site":
                for s in re.split(r"[;|\n]+", v):
                    if s.strip() and s.strip() not in g["sites"]:
                        g["sites"].append(s.strip())
            elif v and not g["data"].get(t):
                g["data"][t] = v

    for g in groups.values():
        data = g["data"]
        existing = None
        if data.get("code"):
            existing = ctx.db.one("SELECT id FROM customers WHERE code = ?", (data["code"],))
        else:
            for r in ctx.db.query("SELECT id, phone FROM customers WHERE norm(name) = ?", (norm(data["name"]),)):
                if digits(r["phone"]) == digits(data.get("phone")):
                    existing = r
                    break
        try:
            if existing:
                if update_existing:
                    customers.update(ctx, existing["id"], data)
                    stats["updated"] += 1
                else:
                    stats["skipped"] += 1
                cid = existing["id"]
                for s in g["sites"]:
                    if not ctx.db.scalar("SELECT 1 FROM sites WHERE customer_id = ? AND name = ?", (cid, s)):
                        customers.add_site(ctx, cid, s)
                        stats["sites"] += 1
            else:
                customers.create(ctx, data, sites=g["sites"])
                stats["created"] += 1
                stats["sites"] += len(g["sites"])
        except ValidationError as e:
            stats["errors"].append((data.get("name"), e.errors[0]))
    with ctx.db.conn:
        audit.log(ctx, "import", "customer", None, {k: v for k, v in stats.items() if k != "errors"})
    return stats


# --- historical tickets ------------------------------------------------------------------

TICKET_TARGETS = ["ticket_no", "date", "time", "customer", "customer_code", "site", "grade", "quantity",
                  "unit_price", "total", "truck_no", "driver", "pump", "pump_no", "plant", "status", "notes"]

TICKET_SYNONYMS = {
    "ticket_no": ["ticket", "ticket no", "ticket number", "bon", "bon no", "no", "number", "delivery note", "dn",
                  "رقم البون", "البون", "بون", "رقم", "رقم الفاتوره", "رقم البوليصه"],
    "date": ["date", "day", "delivery date", "التاريخ", "تاريخ", "تاريخ التسليم"],
    "time": ["time", "hour", "الوقت", "الساعه", "وقت"],
    "customer": ["customer", "client", "customer name", "client name", "name", "الزبون", "اسم الزبون", "العميل",
                 "اسم العميل", "الاسم"],
    "customer_code": ["customer code", "client code", "code", "كود الزبون", "رمز الزبون", "الكود", "كود"],
    "site": SYNONYMS["site"],
    "grade": ["grade", "class", "concrete", "concrete grade", "type", "mix", "الصنف", "صنف", "العيار",
              "صنف الباطون", "نوع الباطون", "الباطون"],
    "quantity": ["quantity", "qty", "m3", "m³", "volume", "cbm", "الكميه", "كميه", "م3", "متر", "الكميه م3"],
    "unit_price": ["price", "unit price", "rate", "price/m3", "السعر", "سعر المتر", "سعر"],
    "total": ["total", "amount", "value", "المجموع", "المبلغ", "القيمه", "الاجمالي", "المجموع $"],
    "truck_no": ["truck", "truck no", "mixer", "mixer no", "plate", "الخلاطه", "رقم الخلاطه", "الشاحنه", "خلاطه"],
    "driver": ["driver", "السائق", "سائق", "الشوفير"],
    "pump": ["pump", "pump used", "مضخه", "المضخه"],
    "pump_no": ["pump no", "pump number", "رقم المضخه"],
    "plant": ["plant", "batching plant", "station", "المجبل", "مجبل", "المحطه"],
    "status": ["status", "الحاله", "حاله"],
    "notes": SYNONYMS["notes"],
}

_YES = {"yes", "y", "1", "x", "✓", "✔", "true", "نعم", "اي", "ايه", "oui", "pump"}
_CANCELLED = {"cancelled", "canceled", "cancel", "void", "annule", "annulé", "ملغي", "ملغى", "ملغا", "الغاء",
              "ملغيه"}


def guess_ticket_mapping(header: list[str]) -> dict:
    return guess_mapping(header, TICKET_SYNONYMS)


def parse_date(text: str, day_first: bool = True) -> str | None:
    """'2026-10-05', '05/10/2026', '5-10-26', '05.10.2026' or an Excel serial number → ISO date."""
    s = norm(text)
    if not s:
        return None
    try:
        m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
        if m:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})\b", s)
        if m:
            a, b, y = int(m[1]), int(m[2]), int(m[3])
            y += 2000 if y < 100 else 0
            d, mo = (a, b) if day_first else (b, a)
            if mo > 12 and d <= 12:  # clearly the other order
                d, mo = mo, d
            return date(y, mo, d).isoformat()
        if re.fullmatch(r"\d{5}(\.\d+)?", s) and 20000 < float(s) < 80000:  # Excel serial day
            return (date(1899, 12, 30) + timedelta(days=int(float(s)))).isoformat()
    except ValueError:
        return None
    return None


def parse_time(text: str) -> str:
    s = norm(text).replace(".", ":")
    m = re.search(r"(\d{1,2}):(\d{2})\s*(am|pm|ص|م)?", s)
    if not m:
        return ""
    h, mi = int(m[1]), int(m[2])
    if m[3] in ("pm", "م") and h < 12:
        h += 12
    if m[3] in ("am", "ص") and h == 12:
        h = 0
    return f"{h:02d}:{mi:02d}" if h < 24 and mi < 60 else ""


def parse_number(text):
    """'1,250.50', '$70', '7,5', '١٢' → float; '' → None; garbage → 'invalid'."""
    s = re.sub(r"[^\d.,\-]", "", norm(text))
    if not s:
        return None if not norm(text) else "invalid"
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif s.count(",") == 1 and len(s.split(",")[1]) != 3:
        s = s.replace(",", ".")  # decimal comma
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return "invalid"


def normalize_grade(text: str) -> str:
    g = re.sub(r"\s+", "", (text or "").strip()).upper()
    return f"B{g}" if g.isdigit() else g


def import_tickets(ctx, rows: list[list[str]], mapping: dict, day_first: bool = True,
                   create_customers: bool = True, dry_run: bool = False) -> dict:
    """Import historical tickets in one transaction.

    Customers are matched by code, then by name (Arabic letter forms ignored); missing ones are created
    when `create_customers` is set. Sites are created under their customer. Ticket numbers that already
    exist, or repeat inside the file, are reported as duplicates and skipped. Rows with errors are
    skipped and reported with their line number. With `dry_run` nothing is saved (a check before import).
    """
    from .services import customers, lookups, tickets

    ctx.require_admin()
    required = {"ticket_no", "date", "quantity"}
    if not required <= set(mapping) or not ({"customer", "customer_code"} & set(mapping)):
        raise ValidationError(["err_map_tickets"])
    stats = {"created": 0, "customers_created": 0, "sites_created": 0, "duplicates": [], "errors": [],
             "quantity": 0.0, "amount": 0.0, "date_from": None, "date_to": None}
    db = ctx.db
    by_code = {norm(r["code"]): r["id"] for r in db.query("SELECT id, code FROM customers WHERE deleted_at IS NULL")}
    by_name = {}
    for r in db.query("SELECT id, name FROM customers WHERE deleted_at IS NULL ORDER BY id"):
        by_name.setdefault(norm(r["name"]), r["id"])
    site_cache = {}
    seen = set()

    def cell(row, key):
        i = mapping.get(key)
        return (row[i] if i is not None and i < len(row) else "").strip()

    try:
        for line, row in enumerate(rows, start=2):
            if not any(c.strip() for c in row):
                continue
            no = cell(row, "ticket_no")
            if no.endswith(".0") and no[:-2].isdigit():
                no = no[:-2]
            if no and (no in seen or tickets.is_duplicate(db, no)):
                stats["duplicates"].append((line, no))
                continue

            # customer
            code, name = cell(row, "customer_code"), cell(row, "customer")
            cid = by_code.get(norm(code)) if code else None
            if cid is None and name:
                cid = by_name.get(norm(name))
            if cid is None:
                if not (create_customers and name):
                    stats["errors"].append((line, [("err_customer_not_found", {"name": name or code})]))
                    continue
                taken = code and db.scalar("SELECT 1 FROM customers WHERE code = ?", (code,))  # incl. Trash
                new_code = code if code and not taken else customers.next_code(db)
                cid = customers.insert(ctx, {"code": new_code, "name": name}, log=False)
                by_code[norm(new_code)] = cid
                by_name[norm(name)] = cid
                stats["customers_created"] += 1

            raw_date = cell(row, "date")
            iso = parse_date(raw_date, day_first)
            t = parse_time(cell(row, "time")) or (parse_time(raw_date[10:]) if len(raw_date) > 10 else "")
            qty = parse_number(cell(row, "quantity"))
            price = parse_number(cell(row, "unit_price"))
            total = parse_number(cell(row, "total"))
            pump_no = cell(row, "pump_no")
            pump_cell = cell(row, "pump")
            pump_used = bool(pump_no) or norm(pump_cell) in _YES or bool(re.search(r"\d", pump_cell))
            if pump_cell and not pump_no and re.search(r"\d", pump_cell):
                pump_no = pump_cell
            if price is None and isinstance(total, float) and isinstance(qty, float) and qty > 0:
                price = round(total / qty, 4)
            d = tickets.normalize({
                "ticket_no": no, "date": iso or "", "time": t, "customer_id": cid,
                "grade": normalize_grade(cell(row, "grade")), "quantity": qty, "unit_price": price,
                "truck_no": cell(row, "truck_no"), "driver": cell(row, "driver"), "pump_used": pump_used,
                "pump_no": pump_no, "plant": cell(row, "plant"), "notes": cell(row, "notes"),
                "status": "cancelled" if norm(cell(row, "status")) in _CANCELLED else "delivered",
            })
            if isinstance(total, float) and isinstance(d["quantity"], float):
                d["total"] = round(total, 2)  # keep the historical amount exactly
            errors = tickets.validate(db, d)
            if total == "invalid":
                errors.append("err_price_invalid")
            if errors:
                stats["errors"].append((line, errors))
                continue

            site = cell(row, "site")
            if site:
                key = (cid, norm(site))
                if key not in site_cache:
                    existing = db.scalar("SELECT id FROM sites WHERE customer_id = ? AND norm(name) = ? "
                                         "AND deleted_at IS NULL", (cid, norm(site)))
                    if existing is None:
                        existing = customers._insert_site(ctx, cid, site)
                        stats["sites_created"] += 1
                    site_cache[key] = existing
                d["site_id"] = site_cache[key]
            tickets.insert(ctx, d, log=False)
            lookups.learn(db, "grade", d["grade"])
            lookups.learn(db, "plant", d["plant"])
            seen.add(no)
            stats["created"] += 1
            if d["status"] == "delivered":
                stats["quantity"] += d["quantity"]
                stats["amount"] += d["total"]
            stats["date_from"] = min(filter(None, [stats["date_from"], d["date"]]))
            stats["date_to"] = max(filter(None, [stats["date_to"], d["date"]]))
        stats["quantity"] = round(stats["quantity"], 2)
        stats["amount"] = round(stats["amount"], 2)
        if dry_run:
            db.conn.rollback()
        else:
            audit.log(ctx, "import", "ticket", None,
                      {k: v for k, v in stats.items() if k not in ("errors", "duplicates")}
                      | {"duplicates": len(stats["duplicates"]), "errors": len(stats["errors"])})
            db.conn.commit()
    except Exception:
        db.conn.rollback()
        raise
    return stats


def write_template(path, kind: str):
    """An empty Excel sheet with the expected headers, for people preparing their old data."""
    from .exporter import export_xlsx
    from .i18n import is_rtl, tr
    if kind == "tickets":
        keys = ["ticket_no", "date", "time", "customer", "code", "site", "grade", "quantity", "unit_price", "total",
                "truck", "driver", "pump_no", "plant", "status", "notes"]
        sample = ["250001", "05/10/2026", "08:30", "Rami Saad", "", "Byblos Marina Tower", "B300", 8, 75, 600,
                  "M-12", "Tony Rizk", "", "Plant 1", "", ""]
    else:
        keys = ["code", "name", "company_name", "phone", "phone2", "area", "address", "payment_terms", "site",
                "notes"]
        sample = ["", "Rami Saad", "Saad Construction", "70 118 342", "", "Jbeil", "Jbeil", "Cash",
                  "Byblos Marina Tower", ""]
    path = Path(path)
    export_xlsx(path, tr("import_template"), [tr(k) for k in keys], [sample], rtl=is_rtl())
    # export_xlsx puts a title block on top; importing expects the header in row 1, so drop it
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    ws.delete_rows(1, 3)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = None
    wb.save(path)
    return path
