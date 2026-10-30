"""Export any table to Excel (.xlsx) or PDF, with proper Arabic text in both."""
import sys
from datetime import datetime
from pathlib import Path

FONT_CANDIDATES = [
    # Windows fonts with Arabic glyphs
    r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\segoeui.ttf",
    # Linux (development / CI)
    "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]
BOLD_CANDIDATES = [
    r"C:\Windows\Fonts\tahomabd.ttf", r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\segoeuib.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansArabic-Bold.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]

_fonts = None


def _register_fonts():
    global _fonts
    if _fonts:
        return _fonts
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular = next((p for p in FONT_CANDIDATES if Path(p).exists()), None)
    bold = next((p for p in BOLD_CANDIDATES if Path(p).exists()), None) or regular
    if regular:
        pdfmetrics.registerFont(TTFont("AracoFont", regular))
        pdfmetrics.registerFont(TTFont("AracoFont-Bold", bold))
        _fonts = ("AracoFont", "AracoFont-Bold")
    else:  # no Arabic-capable font found; Latin text still works
        _fonts = ("Helvetica", "Helvetica-Bold")
    return _fonts


def shape(text) -> str:
    """Join Arabic letters and put them in visual order for the PDF canvas."""
    s = "" if text is None else str(text)
    if not any("\u0600" <= ch <= "\u06ff" for ch in s):
        return s
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(s))


def _fmt(v):
    if isinstance(v, float):
        return f"{v:,.2f}"
    return "" if v is None else str(v)


def export_xlsx(path, title: str, headers: list[str], rows: list[list], totals=None, rtl: bool = False,
                subtitle: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Report")[:31].replace("/", "-")
    ws.sheet_view.rightToLeft = rtl
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle or datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F2A44")
    thin = Side(style="thin", color="BBBBBB")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=4, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([v if v is not None else "" for v in row])
    if totals:
        ws.append(totals)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=ws.max_row, column=col)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FCE4CC")
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in r:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
    for col in range(1, len(headers) + 1):
        width = max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(4, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 50)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(rows)}"
    wb.save(path)
    return Path(path)


def export_pdf(path, title: str, headers: list[str], rows: list[list], totals=None, rtl: bool = False,
               subtitle: str = "", company: str = "ARACO Ready Mix"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    page = landscape(A4) if len(headers) > 6 else A4
    doc = SimpleDocTemplate(str(path), pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    align = 2 if rtl else 0  # TA_RIGHT / TA_LEFT
    cell_style = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10, alignment=align)
    head_style = ParagraphStyle("head", parent=cell_style, fontName=bold, textColor=colors.white, alignment=1)
    title_style = ParagraphStyle("title", fontName=bold, fontSize=15, leading=19, alignment=align)
    sub_style = ParagraphStyle("sub", fontName=font, fontSize=9, leading=12, alignment=align,
                               textColor=colors.HexColor("#555555"))

    def p(v, style=cell_style):
        return Paragraph(shape(_fmt(v)).replace("&", "&amp;").replace("<", "&lt;"), style)

    def order(seq):
        return list(reversed(seq)) if rtl else list(seq)

    data = [order([p(h, head_style) for h in headers])]
    data += [order([p(v) for v in row]) for row in rows]
    if totals:
        data.append(order([p(v, ParagraphStyle("tot", parent=cell_style, fontName=bold)) for v in totals]))

    avail = page[0] - 20 * mm
    weights = []
    for i, h in enumerate(headers):
        longest = max([len(_fmt(h))] + [len(_fmt(r[i])) for r in rows[:300]])
        weights.append(min(max(longest, 4), 30))
    widths = order([avail * w / sum(weights) for w in weights])

    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]),
    ]
    if totals:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FCE4CC")))
    table.setStyle(TableStyle(style))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.drawString(10 * mm, 6 * mm, shape(company))
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f"{d.page}")
        canvas.restoreState()

    story = [p(title, title_style), p(subtitle or datetime.now().strftime("%Y-%m-%d %H:%M"), sub_style),
             Spacer(1, 4 * mm), table]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return Path(path)


def open_file(path):
    """Open an exported file with its default program."""
    import os
    import subprocess
    if sys.platform == "win32":
        os.startfile(str(path))  # noqa: S606
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])
